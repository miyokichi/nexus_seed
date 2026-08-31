"""One ``.xlsx`` sheet becomes one Canonical YAML v0.1 document.

The v0.1 converter is deliberately small: it reads a *tabular* sheet with one
header row, and a mapping file says which column carries which Canonical
field.  There is no document understanding, no layout inference and no model
call — a person who can name their columns can ingest their workbook, and a
person who cannot has a mapping file to write rather than a parser to trust.

What one row may produce:

* several entities (one per configured entity block), each with scalar
  properties and an optional unit,
* the relations the mapping declares between those entities,
* provenance: the source file, the sheet, and the cell range the fields for
  that entity were actually read from.

Excel-specific facts stop here.  Everything downstream sees Canonical YAML —
:mod:`nexus_seed.modules.knowledge.canonical` validates the result, so a
converter bug is a load error rather than a corrupt graph.  Nothing in this
module imports the Knowledge Runtime, a Knowledge backend, or Semantica.

Requires ``openpyxl`` (``pip install 'nexus-seed[ingest]'``), imported lazily
so that reading a workbook stays an optional capability.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from ..modules.knowledge.canonical import (
    CanonicalDocument,
    CanonicalEntity,
    CanonicalRelation,
    SourceLocation,
    canonical_from_mapping,
)


class ExcelSourceError(ValueError):
    """Raised when a mapping or a workbook cannot produce Canonical YAML."""


@dataclass(frozen=True, slots=True)
class FieldSpec:
    """Where one Canonical value comes from: a column, or a fixed literal."""

    column: str | None = None
    literal: Any | None = None

    def read(self, row: Mapping[str, Any]) -> Any:
        """Return this field's value for ``row`` (``None`` when the cell is empty)."""

        if self.column is None:
            return self.literal
        if self.column not in row:
            raise ExcelSourceError(f"column {self.column!r} is not in the sheet header")
        return row[self.column]

    @property
    def columns(self) -> tuple[str, ...]:
        """Return the columns this field reads, for provenance."""

        return (self.column,) if self.column else ()

    @classmethod
    def from_value(cls, value: Any, *, field_name: str) -> "FieldSpec":
        """Read ``{"column": ...}``, ``{"value": ...}``, or a bare literal."""

        if isinstance(value, Mapping):
            unknown = sorted(set(value) - {"column", "value"})
            if unknown:
                raise ExcelSourceError(f"{field_name} has unknown keys: {unknown}")
            if ("column" in value) == ("value" in value):
                raise ExcelSourceError(
                    f"{field_name} needs exactly one of 'column' or 'value'"
                )
            if "column" in value:
                return cls(column=_text(value["column"], f"{field_name}.column"))
            return cls(literal=value["value"])
        if isinstance(value, str):
            return cls(literal=value)
        raise ExcelSourceError(f"{field_name} must be a string or a mapping")


@dataclass(frozen=True, slots=True)
class PropertySpec:
    """One Canonical property, optionally carrying its unit of measure."""

    name: str
    value: FieldSpec
    unit: FieldSpec | None = None

    def read(self, row: Mapping[str, Any]) -> Any:
        """Return the property value, or ``{"value": ..., "unit": ...}`` with a unit."""

        value = self.value.read(row)
        if value is None:
            return None
        if self.unit is None:
            return value
        unit = self.unit.read(row)
        return {"value": value, "unit": unit} if unit is not None else value

    @property
    def columns(self) -> tuple[str, ...]:
        unit_columns = self.unit.columns if self.unit else ()
        return (*self.value.columns, *unit_columns)


@dataclass(frozen=True, slots=True)
class EntitySpec:
    """How one row's columns become one Canonical entity."""

    key: str
    id: FieldSpec
    name: FieldSpec
    type: FieldSpec
    aliases: FieldSpec | None = None
    properties: tuple[PropertySpec, ...] = ()

    @property
    def columns(self) -> tuple[str, ...]:
        alias_columns = self.aliases.columns if self.aliases else ()
        return (
            *self.id.columns,
            *self.name.columns,
            *self.type.columns,
            *alias_columns,
            *(column for spec in self.properties for column in spec.columns),
        )


@dataclass(frozen=True, slots=True)
class RelationSpec:
    """One source-asserted relation between two entities of the same row."""

    subject: str
    predicate: FieldSpec
    object: str


@dataclass(frozen=True, slots=True)
class ExcelMapping:
    """The whole column-to-Canonical mapping for one sheet."""

    document_id: str
    sheet: str
    entities: tuple[EntitySpec, ...]
    relations: tuple[RelationSpec, ...] = ()
    header_row: int = 1
    version: str = "0.1"

    @classmethod
    def from_mapping(cls, value: Mapping[str, Any]) -> "ExcelMapping":
        """Validate and normalize a decoded mapping file."""

        if not isinstance(value, Mapping):
            raise ExcelSourceError("mapping root must be a mapping")
        unknown = sorted(
            set(value) - {"version", "document", "sheet", "header_row", "entities", "relations"}
        )
        if unknown:
            raise ExcelSourceError(f"mapping root has unknown keys: {unknown}")
        document = value.get("document")
        if not isinstance(document, Mapping):
            raise ExcelSourceError("mapping.document must be a mapping")
        document_unknown = sorted(set(document) - {"id"})
        if document_unknown:
            raise ExcelSourceError(f"mapping.document has unknown keys: {document_unknown}")
        raw_entities = value.get("entities")
        if not isinstance(raw_entities, list) or not raw_entities:
            raise ExcelSourceError("mapping.entities must be a non-empty list")
        entities = tuple(
            _entity_spec(item, index=index) for index, item in enumerate(raw_entities)
        )
        keys = [entity.key for entity in entities]
        if len(keys) != len(set(keys)):
            raise ExcelSourceError("mapping.entities keys must be unique")
        raw_relations = value.get("relations", [])
        if not isinstance(raw_relations, list):
            raise ExcelSourceError("mapping.relations must be a list")
        relations = tuple(
            _relation_spec(item, index=index, keys=set(keys))
            for index, item in enumerate(raw_relations)
        )
        header_row = value.get("header_row", 1)
        if not isinstance(header_row, int) or isinstance(header_row, bool) or header_row < 1:
            raise ExcelSourceError("mapping.header_row must be a positive integer")
        return cls(
            document_id=_text(document.get("id"), "mapping.document.id"),
            sheet=_text(value.get("sheet"), "mapping.sheet"),
            entities=entities,
            relations=relations,
            header_row=header_row,
            version=str(value.get("version", "0.1")),
        )


def loads_excel_mapping(text: str) -> ExcelMapping:
    """Parse one mapping file from YAML text."""

    decoded = _yaml().safe_load(text)
    return ExcelMapping.from_mapping(decoded)


def load_excel_mapping(path: str | Path) -> ExcelMapping:
    """Read and validate one UTF-8 mapping YAML file."""

    return loads_excel_mapping(Path(path).read_text(encoding="utf-8"))


def convert_excel(
    workbook_path: str | Path,
    mapping: ExcelMapping | str | Path,
    *,
    document_id: str | None = None,
) -> CanonicalDocument:
    """Convert one sheet of ``workbook_path`` into a validated Canonical document."""

    spec = mapping if isinstance(mapping, ExcelMapping) else load_excel_mapping(mapping)
    path = Path(workbook_path)
    rows = _read_rows(path, sheet=spec.sheet, header_row=spec.header_row)
    identifier = document_id or spec.document_id

    entities: dict[str, CanonicalEntity] = {}
    relations: dict[tuple[str, str, str], CanonicalRelation] = {}
    for row_number, row in rows:
        row_ids: dict[str, str] = {}
        for entity_spec in spec.entities:
            entity = _entity(entity_spec, row, row_number, spec.sheet, identifier)
            if entity is None:
                continue
            row_ids[entity_spec.key] = entity.id
            existing = entities.get(entity.id)
            entities[entity.id] = entity if existing is None else _merge(existing, entity)
        for relation_spec in spec.relations:
            subject = row_ids.get(relation_spec.subject)
            obj = row_ids.get(relation_spec.object)
            predicate = _clean(relation_spec.predicate.read(row))
            if subject is None or obj is None or not predicate:
                continue
            key = (subject, str(predicate), obj)
            if key in relations:
                continue
            columns = (
                *_spec_by_key(spec, relation_spec.subject).columns,
                *_spec_by_key(spec, relation_spec.object).columns,
                *relation_spec.predicate.columns,
            )
            relations[key] = CanonicalRelation(
                subject=subject,
                predicate=str(predicate),
                object=obj,
                source=_location(columns, row, row_number, spec.sheet, identifier),
            )

    if not entities:
        raise ExcelSourceError(
            f"{path.name}!{spec.sheet} produced no entities; check the mapping columns"
        )
    document = CanonicalDocument(
        document_id=identifier,
        source_file=path.name,
        source_location=spec.sheet,
        text=_text_view(tuple(entities.values()), tuple(relations.values())),
        entities=tuple(entities.values()),
        explicit_relations=tuple(relations.values()),
    )
    # Round-trip through the Canonical validator: the converter never hands
    # downstream a document the Knowledge boundary would reject.
    return canonical_from_mapping(document.to_dict())


def convert_excel_to_yaml(
    workbook_path: str | Path,
    mapping: ExcelMapping | str | Path,
    *,
    document_id: str | None = None,
) -> str:
    """Return the Canonical YAML v0.1 text for one sheet."""

    document = convert_excel(workbook_path, mapping, document_id=document_id)
    return _yaml().safe_dump(
        document.to_dict(),
        allow_unicode=True,
        sort_keys=False,
        default_flow_style=False,
    )


def _entity(
    spec: EntitySpec,
    row: Mapping[str, Any],
    row_number: int,
    sheet: str,
    document_id: str,
) -> CanonicalEntity | None:
    identifier = _clean(spec.id.read(row))
    if not identifier:
        return None
    name = _clean(spec.name.read(row)) or identifier
    entity_type = _clean(spec.type.read(row))
    if not entity_type:
        raise ExcelSourceError(
            f"entity {spec.key!r} has no type in row {row_number} of {sheet!r}"
        )
    properties = {}
    for property_spec in spec.properties:
        value = property_spec.read(row)
        if value is not None:
            properties[property_spec.name] = value
    aliases = _aliases(spec.aliases.read(row)) if spec.aliases else ()
    return CanonicalEntity(
        id=str(identifier),
        name=str(name),
        type=str(entity_type),
        aliases=aliases,
        properties=properties,
        source=_location(spec.columns, row, row_number, sheet, document_id),
    )


def _merge(existing: CanonicalEntity, addition: CanonicalEntity) -> CanonicalEntity:
    """Keep one entity per id: the first row defines it, later rows fill gaps.

    A workbook repeats the same subject on every row it is relevant to.  The
    first occurrence is the one a reader is sent back to, so identity, type and
    provenance come from it; later rows may only add properties and aliases it
    did not carry.
    """

    properties = dict(addition.properties)
    properties.update(existing.properties)
    aliases = (*existing.aliases, *(a for a in addition.aliases if a not in existing.aliases))
    return CanonicalEntity(
        id=existing.id,
        name=existing.name,
        type=existing.type,
        aliases=aliases,
        properties=properties,
        source=existing.source,
    )


def _location(
    columns: Iterable[str],
    row: Mapping[str, Any],
    row_number: int,
    sheet: str,
    document_id: str,
) -> SourceLocation:
    """Return the ``Sheet!B12:F12`` range the values were actually read from."""

    letters = sorted(
        {
            row[_LETTERS][column]
            for column in columns
            if column and column in row[_LETTERS]
        },
        key=_column_index,
    )
    if not letters:
        return SourceLocation(document_id=document_id, location=sheet)
    first, last = letters[0], letters[-1]
    cells = f"{first}{row_number}" if first == last else f"{first}{row_number}:{last}{row_number}"
    return SourceLocation(document_id=document_id, location=f"{sheet}!{cells}")


def _text_view(
    entities: Sequence[CanonicalEntity], relations: Sequence[CanonicalRelation]
) -> str:
    """Render the same facts as plain sentences, for readers and for retrieval."""

    names = {entity.id: entity.name for entity in entities}
    lines = []
    for entity in entities:
        head = f"{entity.name} ({entity.type})"
        if entity.properties:
            details = ", ".join(
                f"{key} = {_readable(value)}"
                for key, value in sorted(entity.properties.items())
            )
            lines.append(f"{head}: {details}.")
        else:
            lines.append(f"{head}.")
    for relation in relations:
        subject = names.get(relation.subject, relation.subject)
        obj = names.get(relation.object, relation.object)
        lines.append(f"{subject} {relation.predicate.replace('_', ' ')} {obj}.")
    return "\n".join(lines)


def _readable(value: Any) -> str:
    if isinstance(value, Mapping) and set(value) == {"value", "unit"}:
        return f"{_readable(value['value'])} {value['unit']}"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _aliases(value: Any) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, (list, tuple)):
        parts = [str(item) for item in value]
    else:
        parts = str(value).split(",")
    return tuple(dict.fromkeys(part.strip() for part in parts if part.strip()))


#: Key under which each row carries its column-name -> column-letter map.  A
#: private, non-string key so it can never collide with a header cell.
_LETTERS = object()


def _read_rows(
    path: Path, *, sheet: str, header_row: int
) -> list[tuple[int, dict[Any, Any]]]:
    """Read one sheet as ``(row_number, {column: value})`` pairs."""

    openpyxl = _openpyxl()
    if not path.is_file():
        raise ExcelSourceError(f"workbook not found: {path}")
    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many shapes
        raise ExcelSourceError(f"{path.name} is not a readable .xlsx: {exc}") from exc
    try:
        if sheet not in workbook.sheetnames:
            raise ExcelSourceError(
                f"{path.name} has no sheet {sheet!r}; found {workbook.sheetnames}"
            )
        worksheet = workbook[sheet]
        letters: dict[str, str] = {}
        headers: list[str] = []
        rows: list[tuple[int, dict[Any, Any]]] = []
        for row_number, cells in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_number < header_row:
                continue
            if row_number == header_row:
                headers = [
                    "" if cell is None else str(cell).strip() for cell in cells
                ]
                letters = {
                    header: openpyxl.utils.get_column_letter(index)
                    for index, header in enumerate(headers, start=1)
                    if header
                }
                if not letters:
                    raise ExcelSourceError(
                        f"{path.name}!{sheet} row {header_row} has no column headers"
                    )
                continue
            if all(cell is None or str(cell).strip() == "" for cell in cells):
                continue
            # Every header keeps a key even when the row is short, so a
            # trailing empty cell reads as "no value" rather than as a
            # mapping that names a column the sheet does not have.
            row: dict[Any, Any] = {
                header: _cell(cells[index] if index < len(cells) else None)
                for index, header in enumerate(headers)
                if header
            }
            row[_LETTERS] = letters
            rows.append((row_number, row))
        if not rows:
            raise ExcelSourceError(f"{path.name}!{sheet} has no data rows")
        return rows
    finally:
        workbook.close()


def _cell(value: Any) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _column_index(letter: str) -> int:
    index = 0
    for character in letter:
        index = index * 26 + (ord(character.upper()) - ord("A") + 1)
    return index


def _entity_spec(value: Any, *, index: int) -> EntitySpec:
    field_name = f"mapping.entities[{index}]"
    if not isinstance(value, Mapping):
        raise ExcelSourceError(f"{field_name} must be a mapping")
    unknown = sorted(set(value) - {"key", "id", "name", "type", "aliases", "properties"})
    if unknown:
        raise ExcelSourceError(f"{field_name} has unknown keys: {unknown}")
    for required in ("key", "id", "type"):
        if required not in value:
            raise ExcelSourceError(f"{field_name}.{required} is required")
    identifier = FieldSpec.from_value(value["id"], field_name=f"{field_name}.id")
    raw_properties = value.get("properties", {})
    if not isinstance(raw_properties, Mapping):
        raise ExcelSourceError(f"{field_name}.properties must be a mapping")
    return EntitySpec(
        key=_text(value.get("key"), f"{field_name}.key"),
        id=identifier,
        name=(
            FieldSpec.from_value(value["name"], field_name=f"{field_name}.name")
            if "name" in value
            else identifier
        ),
        type=FieldSpec.from_value(value["type"], field_name=f"{field_name}.type"),
        aliases=(
            FieldSpec.from_value(value["aliases"], field_name=f"{field_name}.aliases")
            if "aliases" in value
            else None
        ),
        properties=tuple(
            _property_spec(name, spec, field_name=f"{field_name}.properties")
            for name, spec in raw_properties.items()
        ),
    )


def _property_spec(name: Any, value: Any, *, field_name: str) -> PropertySpec:
    property_name = _text(name, f"{field_name} key")
    if isinstance(value, Mapping) and "unit" in value:
        rest = {key: item for key, item in value.items() if key != "unit"}
        return PropertySpec(
            name=property_name,
            value=FieldSpec.from_value(rest, field_name=f"{field_name}.{property_name}"),
            unit=FieldSpec.from_value(
                value["unit"], field_name=f"{field_name}.{property_name}.unit"
            ),
        )
    return PropertySpec(
        name=property_name,
        value=FieldSpec.from_value(value, field_name=f"{field_name}.{property_name}"),
    )


def _relation_spec(value: Any, *, index: int, keys: set[str]) -> RelationSpec:
    field_name = f"mapping.relations[{index}]"
    if not isinstance(value, Mapping):
        raise ExcelSourceError(f"{field_name} must be a mapping")
    unknown = sorted(set(value) - {"subject", "predicate", "object"})
    if unknown:
        raise ExcelSourceError(f"{field_name} has unknown keys: {unknown}")
    subject = _text(value.get("subject"), f"{field_name}.subject")
    obj = _text(value.get("object"), f"{field_name}.object")
    for key in (subject, obj):
        if key not in keys:
            raise ExcelSourceError(f"{field_name} references unknown entity key {key!r}")
    if "predicate" not in value:
        raise ExcelSourceError(f"{field_name}.predicate is required")
    return RelationSpec(
        subject=subject,
        predicate=FieldSpec.from_value(
            value["predicate"], field_name=f"{field_name}.predicate"
        ),
        object=obj,
    )


def _spec_by_key(mapping: ExcelMapping, key: str) -> EntitySpec:
    for spec in mapping.entities:
        if spec.key == key:
            return spec
    raise ExcelSourceError(f"unknown entity key {key!r}")


def _text(value: Any, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ExcelSourceError(f"{field_name} must be a non-empty string")
    return value.strip()


def _yaml():
    try:
        import yaml
    except ImportError as exc:  # pragma: no cover - packaging guard
        raise ExcelSourceError("PyYAML is required to read a mapping file") from exc
    return yaml


def _openpyxl():
    try:
        import openpyxl
        import openpyxl.utils  # noqa: F401 - ensure the utils namespace is bound
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ExcelSourceError(
            "openpyxl is required to read .xlsx. Install it with:\n"
            "    pip install 'nexus-seed[ingest]'"
        ) from exc
    return openpyxl


__all__ = [
    "EntitySpec",
    "ExcelMapping",
    "ExcelSourceError",
    "FieldSpec",
    "PropertySpec",
    "RelationSpec",
    "convert_excel",
    "convert_excel_to_yaml",
    "load_excel_mapping",
    "loads_excel_mapping",
]
